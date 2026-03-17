from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
for vendor_name in (".vendor_pptx", ".vendor"):
    vendor_path = ROOT / vendor_name
    if vendor_path.exists():
        sys.path.insert(0, str(vendor_path))
        break

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUTFILE = ROOT / "AI_for_Physics_and_Math_Research_ORBIT_Case_Study.pptx"

NAVY = RGBColor(12, 22, 46)
MIDNIGHT = RGBColor(20, 32, 62)
SLATE = RGBColor(34, 47, 84)
TEAL = RGBColor(28, 182, 164)
GOLD = RGBColor(232, 181, 74)
CORAL = RGBColor(234, 111, 92)
PALE = RGBColor(244, 246, 250)
MUTED = RGBColor(180, 190, 210)
WHITE = RGBColor(255, 255, 255)
SUCCESS = RGBColor(93, 197, 137)
WARNING = RGBColor(246, 193, 81)


def prs_setup() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def add_bg(slide, color=NAVY):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

    band = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.28)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = GOLD
    band.line.fill.background()

    accent = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(10.65), Inches(0.28), Inches(2.68), Inches(0.12)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()


def add_footer(slide, page_num: int, subtitle: str = "ORBIT case study"):
    box = slide.shapes.add_textbox(Inches(0.4), Inches(7.05), Inches(12.5), Inches(0.25))
    p = box.text_frame.paragraphs[0]
    p.text = f"{subtitle} | Slide {page_num}"
    p.font.name = "Aptos"
    p.font.size = Pt(10)
    p.font.color.rgb = MUTED
    p.alignment = PP_ALIGN.RIGHT


def add_title(slide, title: str, subtitle: str | None = None):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.62), Inches(12.1), Inches(1.1))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Aptos Display"
    p.font.bold = True
    p.font.size = Pt(28)
    p.font.color.rgb = WHITE
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.name = "Aptos"
        p2.font.size = Pt(13)
        p2.font.color.rgb = MUTED


def add_body_box(
    slide,
    left,
    top,
    width,
    height,
    title=None,
    body=None,
    fill=SLATE,
    title_color=WHITE,
    body_color=PALE,
):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = MIDNIGHT
    shape.line.width = Pt(1.0)

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(10)
    tf.margin_right = Pt(10)
    tf.margin_top = Pt(8)
    tf.margin_bottom = Pt(8)
    tf.clear()

    if title:
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Aptos"
        p.font.bold = True
        p.font.size = Pt(15)
        p.font.color.rgb = title_color
        p.space_after = Pt(4)

    if body:
        lines = body if isinstance(body, list) else [body]
        for idx, line in enumerate(lines):
            p = tf.add_paragraph() if title or idx > 0 else tf.paragraphs[0]
            p.text = line
            p.font.name = "Aptos"
            p.font.size = Pt(11.5)
            p.font.color.rgb = body_color
    return shape


def add_bullets(slide, left, top, width, height, bullets, font_size=16, color=PALE):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(6)
    tf.margin_right = Pt(6)
    tf.margin_top = Pt(2)
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"- {line}"
        p.font.name = "Aptos"
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.space_after = Pt(6)
    return box


def add_quote_card(slide, left, top, width, height, quote, outcome, accent=TEAL):
    base = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    base.fill.solid()
    base.fill.fore_color.rgb = SLATE
    base.line.color.rgb = accent
    base.line.width = Pt(2)

    stripe = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, Inches(0.12), height)
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()

    tf = base.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = f'"{quote}"'
    p.font.name = "Aptos"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = WHITE
    p.space_after = Pt(6)

    p2 = tf.add_paragraph()
    p2.text = outcome
    p2.font.name = "Aptos"
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = PALE
    return base


def add_process_arrow(slide, x1, y1, x2, y2, color=TEAL):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    line.line.color.rgb = color
    line.line.width = Pt(2)
    line.line.end_arrowhead = True
    return line


def add_bar_chart(slide, left, top, width, height, items, title):
    add_body_box(slide, left, top, width, height, title=title, body=[], fill=SLATE)
    chart_left = left + Inches(0.25)
    chart_top = top + Inches(0.62)
    chart_width = width - Inches(0.5)
    max_val = max(v for _, v, _ in items)
    label_width = chart_width * 0.34
    bar_area = chart_width * 0.58
    for idx, (name, value, color) in enumerate(items):
        row_top = chart_top + Inches(0.18 + idx * 0.48)
        label = slide.shapes.add_textbox(chart_left, row_top, label_width, Inches(0.28))
        lp = label.text_frame.paragraphs[0]
        lp.text = name
        lp.font.name = "Aptos"
        lp.font.size = Pt(11)
        lp.font.color.rgb = PALE

        bar_len = max(Inches(0.12), bar_area * (value / max_val))
        bar = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            chart_left + label_width + Inches(0.08),
            row_top + Inches(0.02),
            bar_len,
            Inches(0.18),
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = color
        bar.line.fill.background()

        value_box = slide.shapes.add_textbox(
            chart_left + label_width + bar_area + Inches(0.15), row_top, Inches(0.8), Inches(0.28)
        )
        vp = value_box.text_frame.paragraphs[0]
        vp.text = f"{value:.1f}s"
        vp.font.name = "Aptos"
        vp.font.size = Pt(11)
        vp.font.color.rgb = MUTED


def add_matrix(slide, left, top, width, height, columns, rows, cells):
    col_width = width / len(columns)
    row_height = height / (len(rows) + 1)
    for i, col in enumerate(columns):
        add_body_box(
            slide,
            left + col_width * i,
            top,
            col_width - Inches(0.04),
            row_height - Inches(0.03),
            title=col,
            body=[],
            fill=TEAL if i == 0 else SLATE,
            title_color=NAVY if i == 0 else WHITE,
        )
    for r, row_name in enumerate(rows):
        add_body_box(
            slide,
            left,
            top + row_height * (r + 1),
            col_width - Inches(0.04),
            row_height - Inches(0.03),
            title=row_name,
            body=[],
            fill=MIDNIGHT,
        )
        for c in range(1, len(columns)):
            fill = RGBColor(40, 58, 96) if (r + c) % 2 else RGBColor(31, 47, 84)
            add_body_box(
                slide,
                left + col_width * c,
                top + row_height * (r + 1),
                col_width - Inches(0.04),
                row_height - Inches(0.03),
                body=cells[r][c - 1],
                fill=fill,
            )


def build():
    prs = prs_setup()

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, NAVY)
    title_box = slide.shapes.add_textbox(Inches(0.65), Inches(1.05), Inches(12.0), Inches(1.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "AI as a Research Accelerator in Physics and Mathematics"
    p.font.name = "Aptos Display"
    p.font.bold = True
    p.font.size = Pt(30)
    p.font.color.rgb = WHITE
    p2 = tf.add_paragraph()
    p2.text = (
        "A case study from ORBIT: operator bases, IBP and field-redefinition reduction, "
        "Matchete translation, and Einstein-Hilbert matching"
    )
    p2.font.name = "Aptos"
    p2.font.size = Pt(17)
    p2.font.color.rgb = MUTED
    add_body_box(
        slide,
        Inches(0.78),
        Inches(3.05),
        Inches(5.45),
        Inches(2.6),
        title="Thesis",
        body=[
            "Agentic AI is most valuable when the research task is a long chain of reasoning, code repair, symbolic computation, validation, and reporting.",
            "This ORBIT project is a concrete example: the agent repaired a graviton-basis toolkit, built a comparison package, validated it, accelerated the workflow, and clarified the physics.",
        ],
        fill=SLATE,
    )
    add_body_box(
        slide,
        Inches(6.55),
        Inches(3.05),
        Inches(5.95),
        Inches(2.6),
        title="Core result",
        body=[
            "The Matchete-derived result can be reduced and compared systematically to Einstein-Hilbert plus cosmological constant.",
            "At d <= 5 it matches a sign-flipped EH convention before canonical normalization; at d <= 6 a leftover {2,4} sector obstructs a pure EH+Lambda interpretation.",
        ],
        fill=MIDNIGHT,
        title_color=GOLD,
    )
    add_footer(slide, 1)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(
        slide,
        "1. Mathematical and Physical Background",
        "The ORBIT package works in abstract-index xAct syntax and classifies graviton operators by field number and derivative count.",
    )
    add_bullets(
        slide,
        Inches(0.65),
        Inches(1.55),
        Inches(5.8),
        Inches(4.9),
        [
            "Sector notation: V_{n_h,N_d} collects scalar operators built from n_h gravitons h_{mu nu} and N_d derivatives.",
            "Reduction step 1: quotient by integration by parts (IBP), which removes total derivatives and identifies operators differing by divergences.",
            "Reduction step 2: quotient by the image of field redefinitions, implemented in ORBIT through the linearized Fierz-Pauli equations of motion in interaction sectors.",
            "Quadratic sectors are treated separately: for the kinetic theory one usually keeps the full IBP basis and does not mod out by linear field rescalings.",
        ],
        font_size=16,
    )
    add_body_box(
        slide,
        Inches(6.75),
        Inches(1.65),
        Inches(5.8),
        Inches(1.15),
        title="Key quotient picture",
        body=["Raw operators -> modulo IBP -> modulo field redefinitions -> physical coordinates"],
        fill=TEAL,
        body_color=NAVY,
        title_color=NAVY,
    )
    add_body_box(
        slide,
        Inches(6.75),
        Inches(3.05),
        Inches(5.8),
        Inches(2.5),
        title="Reference action used in the comparison package",
        body=[
            "L_ref = Lambda_ref sqrt(-g) + s_EH (2/kappa_ref^2) sqrt(-g) R",
            "mostly-minus convention, with metric perturbation g = eta + kappa h",
            "Comparison is done sector-by-sector in the same reduced basis as the translated Matchete output.",
        ],
        fill=SLATE,
    )
    add_footer(slide, 2)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "2. Stating the Problem")
    add_body_box(
        slide,
        Inches(0.65),
        Inches(1.45),
        Inches(12.0),
        Inches(1.15),
        title="Research question",
        body=[
            "Can we take a Matchete-generated effective Lagrangian, translate it into xAct, remove IBP and field-redefinition redundancies, construct the Einstein-Hilbert reference through the same operator basis, and determine whether the two match?"
        ],
        fill=TEAL,
        body_color=NAVY,
        title_color=NAVY,
    )
    stages = [
        ("Matchete .mx", "Load the stored symbolic output"),
        ("Translation", "Map Matchete syntax to ORBIT/xAct"),
        ("Normalization", "Set hbar -> 1 and keep the finite epsilon^0 part"),
        ("Reduction", "Remove IBP and field-redefinition redundancies"),
        ("Reference build", "Expand EH + Lambda in the same basis"),
        ("Comparison", "Solve for constraints and isolate leftovers"),
    ]
    x = 0.82
    for idx, (title, body) in enumerate(stages):
        add_body_box(
            slide,
            Inches(x),
            Inches(3.0),
            Inches(1.92),
            Inches(2.15),
            title=title,
            body=[body],
            fill=SLATE if idx % 2 == 0 else MIDNIGHT,
        )
        if idx < len(stages) - 1:
            add_process_arrow(slide, Inches(x + 1.9), Inches(4.07), Inches(x + 2.15), Inches(4.07), GOLD)
        x += 2.08
    add_body_box(
        slide,
        Inches(3.25),
        Inches(5.45),
        Inches(6.7),
        Inches(0.82),
        body=[
            "Success criterion: equality in supported sectors, no unsupported leftovers, and a clear statement of the required parameter constraints and conventions."
        ],
        fill=CORAL,
    )
    add_footer(slide, 3)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "3. The Challenges in Solving This Problem")
    challenge_cards = [
        ("Tensor syntax fragility", "Both kernels initially failed with syntax issues. Small xAct mistakes can invalidate long symbolic runs."),
        ("Basis ambiguity", "Notebook conventions and package conventions differed. Some counts changed depending on whether hh(ddh) was reduced correctly by IBP."),
        ("Quadratic-sector subtlety", "Field redefinitions at quadratic order can quotient out the kinetic term itself unless handled carefully."),
        ("Canonical normalization", "Matching depends on whether one normalizes to +L_FP or allows a sign-flipped kinetic convention."),
        ("Performance", "xPert expansion, ToCanonical, SolveConstants, and basis projection are expensive and mostly single-kernel bottlenecks."),
        ("Scientific traceability", "Every claim needed a reproducible check: saved scripts, cached runs, reports, and consistency tests."),
    ]
    positions = [(0.7, 1.65), (4.55, 1.65), (8.4, 1.65), (0.7, 4.1), (4.55, 4.1), (8.4, 4.1)]
    for (x, y), (title, body) in zip(positions, challenge_cards):
        add_body_box(slide, Inches(x), Inches(y), Inches(3.35), Inches(1.95), title=title, body=[body], fill=SLATE)
    add_footer(slide, 4)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, MIDNIGHT)
    add_title(slide, "4. The Recent Boom in Agentic AI Usage")
    add_body_box(
        slide,
        Inches(0.7),
        Inches(1.55),
        Inches(3.85),
        Inches(3.15),
        title="What changed",
        body=[
            "Recent systems are not just text generators. They search, read, execute code, keep state, and iterate across long tasks.",
            "The useful abstraction is an agent: a model with tools, persistence, and the ability to plan and validate intermediate steps.",
        ],
        fill=SLATE,
    )
    add_body_box(
        slide,
        Inches(4.75),
        Inches(1.55),
        Inches(3.85),
        Inches(3.15),
        title="Capabilities that matter for research",
        body=[
            "Codebase archaeology and repair",
            "Symbolic debugging with external tools",
            "Experiment orchestration and caching",
            "Long-horizon state tracking through context files and result artifacts",
            "Automatic report and presentation generation",
        ],
        fill=TEAL,
        body_color=NAVY,
        title_color=NAVY,
    )
    add_body_box(
        slide,
        Inches(8.8),
        Inches(1.55),
        Inches(3.85),
        Inches(3.15),
        title="External signals",
        body=[
            "OpenAI's deep research workflow explicitly targets multistep internet research and synthesis.",
            "OpenAI's 2025 scientific-research evaluation framed agentic systems as end-to-end scientific assistants, not just answer engines.",
        ],
        fill=SLATE,
    )
    add_body_box(
        slide,
        Inches(0.7),
        Inches(5.0),
        Inches(11.95),
        Inches(1.0),
        title="Primary sources used for this slide",
        body=[
            "OpenAI: Introducing deep research (Feb. 2, 2025). OpenAI: Evaluating AI's ability to perform scientific research tasks (Dec. 16, 2025)."
        ],
        fill=MIDNIGHT,
        title_color=GOLD,
    )
    add_footer(slide, 5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "Physics-Driven AI Research: Why This Matters Here")
    add_body_box(
        slide,
        Inches(0.7),
        Inches(1.55),
        Inches(12.0),
        Inches(0.95),
        body=[
            "This project is not generic AI for coding. It is AI used in a physics workflow: tensor conventions, operator bases, perturbative gravity, canonical normalization, and symbolic equivalence modulo redundancies."
        ],
        fill=TEAL,
        body_color=NAVY,
    )
    timeline = [
        ("Feb. 2025", "Deep research", "General multistep retrieval, synthesis, and tool use."),
        ("Dec. 2025", "Scientific research evaluation", "Benchmarks aimed at long-horizon research tasks."),
        ("Feb. 2026", "GPT-5.2 derives a new result in theoretical physics", "OpenAI highlighted genuine physics reasoning on scattering-amplitude problems."),
        ("Mar. 2026", "Extending single-minus amplitudes to gravitons", "OpenAI released a second QFT-focused result, now explicitly in gravity amplitudes."),
    ]
    y = 2.85
    for idx, (date, title, body) in enumerate(timeline):
        add_body_box(
            slide,
            Inches(0.95 + idx * 3.05),
            Inches(y),
            Inches(2.6),
            Inches(2.4),
            title=f"{date} | {title}",
            body=[body],
            fill=SLATE if idx % 2 == 0 else MIDNIGHT,
        )
        if idx < len(timeline) - 1:
            add_process_arrow(slide, Inches(3.25 + idx * 3.05), Inches(4.02), Inches(3.55 + idx * 3.05), Inches(4.02), GOLD)
    add_body_box(
        slide,
        Inches(0.95),
        Inches(5.55),
        Inches(11.6),
        Inches(0.8),
        body=[
            "Relevance to ORBIT: the same agentic pattern applies to perturbative gravity packages - read code, fix conventions, run the kernels, compare mathematical branches, and preserve an audit trail."
        ],
        fill=CORAL,
    )
    add_footer(slide, 6)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, MIDNIGHT)
    add_title(slide, "5. How We Can Use Agentic AI for This Problem")
    columns = ["Task layer", "What the agent does", "Why it is useful here", "What still needs human judgment"]
    rows = ["Code", "Physics", "Performance", "Reporting"]
    cells = [
        ["Patch syntax, add APIs, create drivers", "The package evolved while the mathematics was being clarified", "Choose the right conventions and what should count as physical"],
        ["Check sector counts, compare notebook and package logic", "The hard part is often conceptual mismatch, not just a bug", "Interpret whether a mismatch is real or only conventional"],
        ["Time runs, cache expensive sectors, isolate bottlenecks", "The symbolic stack is slow and fragile; targeted acceleration matters", "Decide which approximations or branch tests are worth running"],
        ["Generate reports, scripts, and presentation artifacts", "Research claims need to be inspectable and shareable", "Decide what is publishable, provisional, or stale"],
    ]
    add_matrix(slide, Inches(0.55), Inches(1.45), Inches(12.1), Inches(4.95), columns, rows, cells)
    add_footer(slide, 7)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "6. Case Study Part I: Repair the ORBIT Core")
    add_quote_card(
        slide,
        Inches(0.7),
        Inches(1.65),
        Inches(3.85),
        Inches(1.7),
        "debug graviton_basis_toolkit.wl until it works and gives the expected output",
        "The agent fixed syntax issues, corrected derivative-partition logic, and made the basis-computation pipeline run again.",
    )
    add_quote_card(
        slide,
        Inches(4.75),
        Inches(1.65),
        Inches(3.85),
        Inches(1.7),
        "doesn't fieldRedefinitionsAndIBP.nb provide different results?",
        "This forced a careful comparison between notebook conventions and the package's full sector generation. The discrepancy exposed an IBP canonicalization bug.",
    )
    add_quote_card(
        slide,
        Inches(8.8),
        Inches(1.65),
        Inches(3.85),
        Inches(1.7),
        "the single graviton, 2 derivative terms have a physical count of 2, but isn't that a total derivative?",
        "The agent traced this to incomplete scalar-derivative canonicalization and fixed the {1,2} sector so it vanished modulo IBP as it should.",
    )
    add_body_box(
        slide,
        Inches(0.7),
        Inches(3.75),
        Inches(12.0),
        Inches(2.15),
        title="Successful physics-guided code changes",
        body=[
            "Restored scalar-derivative commutation and derivative sorting in the IBP path.",
            "Corrected the interpretation of SolveConstants[0 == 0] so trivial sectors do not masquerade as nontrivial.",
            "Changed the quadratic-sector policy so field-redefinition quotienting is skipped for n_h < 3.",
            "Validated corrected counts such as {3,2}: IBPCount = 14, RedefRank = 4, PhysicalCount = 10.",
        ],
        fill=SLATE,
    )
    add_footer(slide, 8)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "Case Study Part II: Build a General Reduction and Validation Workflow")
    add_quote_card(
        slide,
        Inches(0.7),
        Inches(1.65),
        Inches(3.85),
        Inches(1.7),
        "write a new code that takes a lagrangian and removes the ibp and field-redefinition redundancies",
        "The agent added a reducer API that splits an input by sector, projects it onto the raw basis, removes IBP, and then quotients by redefinition images where appropriate.",
    )
    add_quote_card(
        slide,
        Inches(4.75),
        Inches(1.65),
        Inches(3.85),
        Inches(1.7),
        "test the new code by taking a lagrangian which is the EOM times some general lagrangian",
        "The agent built disguised EOM-image tests in {3,2} and {3,4}; both reduced exactly to zero, confirming the implementation.",
    )
    add_quote_card(
        slide,
        Inches(8.8),
        Inches(1.65),
        Inches(3.85),
        Inches(1.7),
        "create a comprehensive report ... mathematical background, validation, and a worked example",
        "The agent generated notebooks, PDFs, and a paper-style LaTeX report documenting the toolkit, algorithms, and validation logic.",
    )
    add_body_box(
        slide,
        Inches(0.7),
        Inches(3.75),
        Inches(12.0),
        Inches(2.15),
        title="Concrete validation outputs",
        body=[
            "Reduction API exercised on mixed sectors and on redundant EOM images.",
            "Persistent validation scripts, PDF reports, and context tracking were added so crashes would not erase state.",
            "The translated LEFT_only_uv_2_eft_6 xAct input reduced from 1725 raw terms to 202 reduced terms across 10 surviving sectors.",
        ],
        fill=SLATE,
    )
    add_footer(slide, 9)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, MIDNIGHT)
    add_title(slide, "Case Study Part III: ORBIT_BUILD and the Fixed-EH Comparison")
    add_quote_card(
        slide,
        Inches(0.7),
        Inches(1.62),
        Inches(3.85),
        Inches(1.78),
        "continue according to the ORBIT_BUILD fixed-EH comparison plan",
        "The agent vendored the translator, built the comparison package, added staged checkpointing, and created a live ORBIT_BUILD context file.",
    )
    add_quote_card(
        slide,
        Inches(4.75),
        Inches(1.62),
        Inches(3.85),
        Inches(1.78),
        "let's speed up the computation as much as possible",
        "The agent identified xPert expansion and basis projection as the bottlenecks, added disk caching for the reference expansion, and removed redundant canonicalization.",
    )
    add_quote_card(
        slide,
        Inches(8.8),
        Inches(1.62),
        Inches(3.85),
        Inches(1.78),
        "did you consider canonical normalization?",
        "The agent derived the exact kinetic prefactor Z_h, separated raw matching from canonically normalized matching, and exposed a genuine sign-convention branch.",
    )
    add_body_box(
        slide,
        Inches(0.7),
        Inches(3.75),
        Inches(12.0),
        Inches(2.22),
        title="Final branch analysis",
        body=[
            "Before canonical normalization, the d <= 5 comparison matches a sign-flipped EH convention: s_EH = -1, Lambda_ref = 0, log(mu^2/M^2) = -3/2, M^2 kappa^2 = 48, and kappa_ref = kappa.",
            "At full d <= 6, the comparison still leaves an unsupported {2,4} sector, so the result is not pure EH + Lambda.",
            "With the current canonical-normalization step enabled, the sign-flipped branch no longer matches because the input is normalized to +L_FP while the reference quadratic sector is -L_FP.",
        ],
        fill=SLATE,
    )
    add_footer(slide, 10)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "Measured Runtime and Bottlenecks")
    full_items = [
        ("Translation", 310.4, TEAL),
        ("Finite-part normalization", 306.5, TEAL),
        ("Reference expansion", 1466.9, GOLD),
        ("Input reduction", 1608.5, CORAL),
        ("Reference reduction", 675.8, GOLD),
        ("Comparison assembly", 546.6, TEAL),
    ]
    add_bar_chart(slide, Inches(0.7), Inches(1.55), Inches(7.2), Inches(4.85), full_items, "Measured full d <= 6 run: 4915 s total (~81.9 min)")
    add_body_box(
        slide,
        Inches(8.2),
        Inches(1.6),
        Inches(4.45),
        Inches(1.45),
        title="Why it took so long",
        body=[
            "The dominant costs were xPert expansion, ToCanonical, SolveConstants, and projection onto large sector bases. xAct was effectively single-kernel in this workflow."
        ],
        fill=SLATE,
    )
    add_body_box(
        slide,
        Inches(8.2),
        Inches(3.25),
        Inches(4.45),
        Inches(1.3),
        title="After caching and scope reduction",
        body=[
            "The corrected d <= 5 rerun dropped to ~374 s (~6.2 min). The EHSign = -1 sign check without canonical normalization completed in ~8.3 min wall-clock."
        ],
        fill=MIDNIGHT,
        title_color=GOLD,
    )
    add_body_box(
        slide,
        Inches(8.2),
        Inches(4.8),
        Inches(4.45),
        Inches(1.2),
        title="AI contribution",
        body=[
            "The agent did not make xAct parallel, but it did make the workflow practical by profiling, caching, checkpointing, and narrowing expensive branches."
        ],
        fill=TEAL,
        body_color=NAVY,
        title_color=NAVY,
    )
    add_footer(slide, 11)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "Tests, Sanity Checks, and Scientific Guardrails")
    checks = [
        ("Sector sanity", "Verified that the {1,2} sector vanishes modulo IBP after the canonicalization fix.", SUCCESS),
        ("Notebook parity", "Resolved the cubic {3,2} mismatch by identifying the missing IBP reduction and clarifying restricted versus full-sector counting.", SUCCESS),
        ("Redundant images", "Tested L = EOM . delta h with disguised expressions; the reducer returned zero.", SUCCESS),
        ("Kinetic normalization", "Showed the Matchete {2,2} sector is exactly Z_h times the Fierz-Pauli kinetic term.", SUCCESS),
        ("Reference consistency", "Checked that the EH {2,2} sector equals the toolkit's Fierz-Pauli Lagrangian in the same reduced basis.", SUCCESS),
        ("Convention split", "Separated raw sign-convention matching from canonical-normalization matching; identified the unresolved kinetic-sign choice.", WARNING),
        ("Unsupported leftovers", "At d <= 6 the full run still leaves a non-EH {2,4} sector.", WARNING),
    ]
    y = 1.6
    for title, body, color in checks:
        add_quote_card(slide, Inches(0.8), Inches(y), Inches(11.9), Inches(0.65), title, body, accent=color)
        y += 0.75
    add_footer(slide, 12)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, MIDNIGHT)
    add_title(slide, "What the AI Actually Added to the Project")
    add_body_box(
        slide,
        Inches(0.7),
        Inches(1.55),
        Inches(4.0),
        Inches(4.9),
        title="New package capabilities",
        body=[
            "Matchete .mx -> xAct translator vendored into ORBIT",
            "General Lagrangian reducer modulo IBP and field redefinitions",
            "Fixed-EH comparison package with report generation",
            "Explicit EHSign option in the reference builder",
        ],
        fill=SLATE,
    )
    add_body_box(
        slide,
        Inches(4.85),
        Inches(1.55),
        Inches(4.0),
        Inches(4.9),
        title="Successful implementation ideas proposed by the agent",
        body=[
            "Skip field-redefinition quotienting in quadratic sectors",
            "Project to the raw basis first, then quotient by stored IBP relations",
            "Use exact epsilon^0 extraction rather than heuristic pole filtering",
            "Separate supported and unsupported sector comparisons",
            "Maintain a persistent ORBIT_BUILD context file",
        ],
        fill=TEAL,
        body_color=NAVY,
        title_color=NAVY,
    )
    add_body_box(
        slide,
        Inches(9.0),
        Inches(1.55),
        Inches(3.65),
        Inches(4.9),
        title="Research outputs",
        body=[
            "Repaired toolkit and validated sector counts",
            "Generated PDFs, notebooks, reports, and rerun scripts",
            "Produced exact sign-branch solutions for the d <= 5 comparison",
            "Left a reproducible artifact trail rather than only ad hoc chat output",
        ],
        fill=SLATE,
    )
    add_footer(slide, 13)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "Current Physics Result")
    add_body_box(
        slide,
        Inches(0.7),
        Inches(1.55),
        Inches(12.0),
        Inches(1.15),
        title="Most concise statement",
        body=[
            "The Matchete result is not simply right or wrong. It matches one convention branch and fails another."
        ],
        fill=TEAL,
        body_color=NAVY,
        title_color=NAVY,
    )
    add_body_box(
        slide,
        Inches(0.7),
        Inches(2.95),
        Inches(5.75),
        Inches(2.65),
        title="At d <= 5",
        body=[
            "With EHSign = -1 and without canonical normalization: exact supported-sector match.",
            "Required conditions: Lambda_ref = 0, log(mu^2/M^2) = -3/2, M^2 kappa^2 = 48, and kappa_ref = kappa.",
        ],
        fill=SLATE,
    )
    add_body_box(
        slide,
        Inches(6.7),
        Inches(2.95),
        Inches(6.0),
        Inches(2.65),
        title="At d <= 6 and under the current canonical convention",
        body=[
            "A leftover {2,4} sector remains, so the result is not pure EH + Lambda.",
            "If the input is canonically normalized to +L_FP, the sign-flipped EH branch no longer matches. This is the remaining convention problem to resolve.",
        ],
        fill=MIDNIGHT,
        title_color=GOLD,
    )
    add_body_box(
        slide,
        Inches(1.9),
        Inches(5.95),
        Inches(9.45),
        Inches(0.72),
        body=[
            "Interpretation: AI helped turn a vague mismatch into a precise, branch-dependent scientific statement with explicit equations, exact tests, and reusable code."
        ],
        fill=CORAL,
    )
    add_footer(slide, 14)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_title(slide, "7. Next Steps")
    add_bullets(
        slide,
        Inches(0.75),
        Inches(1.55),
        Inches(5.85),
        Inches(4.9),
        [
            "Add a kinetic-sign-aware canonical-normalization stage so +L_FP and -L_FP branches can be tested on equal footing.",
            "Extend the field-redefinition layer beyond the linearized Fierz-Pauli image to action-dependent nonlinear redefinitions.",
            "Add fixed-dimension identities and eventually a true evanescent-operator layer.",
            "Add gauge-invariance filters for sectors such as curvature-squared.",
        ],
        font_size=16,
    )
    add_bullets(
        slide,
        Inches(6.75),
        Inches(1.55),
        Inches(5.85),
        Inches(4.9),
        [
            "Parallelize the outer sector workflow more safely or move the heavy algebra to a more parallel backend.",
            "Regenerate the final comparison report once stale artifacts are retired.",
            "Add CI-style regression tests for key sectors and branch conventions.",
            "Use the same agentic workflow on other EFT-to-reference matching problems.",
        ],
        font_size=16,
    )
    add_body_box(
        slide,
        Inches(0.85),
        Inches(5.8),
        Inches(11.75),
        Inches(0.82),
        title="Suggested framing for the talk",
        body=[
            "Present AI not as an oracle, but as a force multiplier for disciplined research: it accelerates debugging, experimentation, and documentation, while the human still decides the physics."
        ],
        fill=TEAL,
        body_color=NAVY,
        title_color=NAVY,
    )
    add_footer(slide, 15)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, MIDNIGHT)
    add_title(slide, "Sources and Artifacts")
    add_body_box(
        slide,
        Inches(0.7),
        Inches(1.55),
        Inches(5.85),
        Inches(4.95),
        title="Primary external sources",
        body=[
            "OpenAI, Introducing deep research (Feb. 2, 2025)",
            "OpenAI, Evaluating AI's ability to perform scientific research tasks (Dec. 16, 2025)",
            "OpenAI, GPT-5.2 derives a new result in theoretical physics (Feb. 13, 2026)",
            "OpenAI, Extending single-minus amplitudes to gravitons (Mar. 4, 2026)",
        ],
        fill=SLATE,
    )
    add_body_box(
        slide,
        Inches(6.8),
        Inches(1.55),
        Inches(5.85),
        Inches(4.95),
        title="Local ORBIT artifacts used in this talk",
        body=[
            "graviton_basis_toolkit.wl",
            "orbit_fixed_eh_comparison.wl",
            "ORBIT_BUILD_CONTEXT.md",
            "LEFT_only_uv_2_eft_6_xact_reduction_summary.wl",
            "orbit_fixed_eh_comparison_d5_v3_result.wl",
            "orbit_fixed_eh_comparison_d5_free_ref_minus_eh_nocanon_result.wl",
        ],
        fill=SLATE,
    )
    link_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.05), Inches(11.8), Inches(0.6))
    p = link_box.text_frame.paragraphs[0]
    p.text = "Deck generated from build_ai_physics_research_presentation.py so the content can be revised and rebuilt as the package evolves."
    p.font.name = "Aptos"
    p.font.size = Pt(12)
    p.font.color.rgb = MUTED
    p.alignment = PP_ALIGN.CENTER
    add_footer(slide, 16)

    prs.save(OUTFILE)
    print(f"Wrote {OUTFILE}")


if __name__ == "__main__":
    build()
